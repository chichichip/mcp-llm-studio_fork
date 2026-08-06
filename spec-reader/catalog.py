# -*- coding: utf-8 -*-
"""엑셀 표준품 목록 조회 + 볼트↔너트 카테고리 매핑.

역할 분담 (헷갈리지 말 것):
    엑셀   = 색인. 어떤 계열(도면번호)이 있는지 찾는 용도.
    스펙   = 치수 판단. dash 확정은 반드시 read_spec.py 판독 결과로 한다.

VLM 을 쓰지 않으므로 외부망에서 개발·테스트할 수 있다(verify.py 와 같은 취지).

엑셀 컬럼 (문서 기준):
    대분류 | 중분류 | 도면번호 | 부품번호 | 품명 | 소재 | 비고
    한 행 = 부품번호 단위(dash 포함). 품명에 직경이 들어 있다:
    "BOLT, DOUBLE HEX, 0.164" -> 중분류 + 직경
"""

import csv
import os
import re
from collections import OrderedDict

# ---------------------------------------------------------------- 정규화


def norm(s):
    """카테고리/품명 **비교 전용** 정규화 키를 만든다.

    엑셀 값에는 오타와 공백 흔들림이 섞여 있다 — 실제로 확인된 것만 해도
    'CLAMP<,SADDLE'(꺾쇠), 'GASCKET'(철자), 'DESIGNSTANDARD'(공백 누락),
    'PLUG& CAP'·'WASHER& SEAL RING'(& 앞 공백 없음), 'RING,RETAINING, SPIRAL'.

    그래서 **공백을 전부 없앤 키**로 비교한다. 목록에 공백만으로 구분되는 서로 다른
    중분류는 없으므로 안전하고, 위 흔들림이 한 번에 흡수된다. 철자 오타(GASCKET)만은
    공백 규칙으로 못 잡으니 KNOWN_MIDS 에 양쪽 철자를 같이 적어 둔다.

    ⚠ 이 함수의 결과는 **비교에만** 쓴다. 사람에게 보여줄 때는 엑셀에 적힌 원본
    문자열을 그대로 쓴다(출력이 엑셀과 달라 보이면 대조가 안 된다).
    """
    if s is None:
        return ""
    t = str(s).upper()
    t = t.replace("<", "").replace(">", "")      # 'CLAMP<,SADDLE' 류 오타
    t = re.sub(r"\s+", "", t)                    # 공백 전부 제거 (흔들림 흡수)
    return t.strip(",")


# ---------------------------------------------------------------- 알려진 중분류
# 사용자가 확인해 준 실제 엑셀 값 목록. **진실의 원천이 아니라 체크리스트**다.
# 엑셀에 여기 없는 값이 나오면 validate_catalog 가 보고한다(조용히 넘어가면
# 매핑에서 누락돼 '짝 없음'으로 잘못 결론난다).
#
# ⚠ 중분류 이름은 유일하지 않다. 'NUT' 은 대분류 NUT 밑에도, FITTING 밑에도 있고
#   'NUT, COUPLING' 은 TUBING 의 중분류다. 그래서 조회는 항상 (대분류, 중분류) 쌍으로 한다.
KNOWN_MIDS = [
    # 체결류
    "BOLT, HEX", "BOLT, DOUBLE HEX", "BOLT, T-HEAD",
    "STUD, RING LOCKED, SERRATED", "STUD, KEY LOCKED", "STUD",
    "SCREW",                       # 대분류이자 중분류 자리에 그대로 오는 값
    "NUT, PLAIN, HEX", "NUT, SELF-LOCKING, HEX", "NUT, SELF-LOCKING, DOUBLE HEX",
    "NUT, SHANK", "NUT, FLOATING", "NUT, LUG", "NUT, CASTELLATED", "NUT",
    "WASHER, FLAT", "WASHER, LOCK", "WASHER",
    "INSERT, HELICAL COIL", "INSERT, RING LOCKED", "INSERT, KEY LOCKED", "INSERT",
    "RING, LOCK, SERRATED",
    "PIN, HEADLESS", "PIN, HEADED", "PIN, COTTER", "PIN, SPRING",
    "RIVET, SOLID", "RIVET, TUBULAR",
    "RING, RETAINING", "RING, RETAINING, SPIRAL",
    "O-RING", "O-RING, METAL", "SEAL",
    "BEARING, SPHERICAL", "BEARING, SLEEVE", "BEARING, ROD END",
    # 배관/구조
    "NUT, COUPLING", "FERRULE/SLEEVE", "GASKET", "GASCKET",   # GASCKET 은 철자 오타라 양쪽 등록
    "CLAMP, LOOP", "CLAMP, LOOP, CUSHIONED", "CLAMP, LOOP, MULTI TUBE",
    "CLAMP, SADDLE", "CLIP", "BRACKET",
    "ADAPTER", "ELBOW", "TEE", "CROSS", "PLUG & CAP", "WASHER & SEAL RING",
    "DESIGN STANDARD FOR FITTING",
    "COUPLING, V-RETAINER", "FLANGE, V-RETAINER", "COUPLING",
    "SEAL, V-RETAINER COUPLING", "DESIGN STANDARD FOR COUPLING, V-RETAINER",
    "COUPLING, V-BAND", "FLANGE, V-BAND COUPLING",
    "DESIGN STANDARD FOR COUPLING, V-BAND",
    "DESIGN STANDARD FOR HOSE ASSEMBLY",
]
_KNOWN = {norm(m) for m in KNOWN_MIDS}


# ---------------------------------------------------------------- 볼트 ↔ 너트 매핑
# **단순 문자열 치환(BOLT->NUT)으로는 안 된다.** 실제 목록을 보면:
#   BOLT, DOUBLE HEX -> 'NUT, DOUBLE HEX' 는 없고 'NUT, SELF-LOCKING, DOUBLE HEX' 뿐
#   BOLT, HEX        -> PLAIN 과 SELF-LOCKING 둘 다 존재
# 그래서 손으로 쓴 매핑표를 쓰고, 후보가 여럿이면 **되묻는다**.
# 셀프락킹 여부는 설계 요구사항이지 부품 목록에서 유도할 수 있는 값이 아니다.
NUT_FOR_BOLT = OrderedDict([
    ("BOLT, DOUBLE HEX", ["NUT, SELF-LOCKING, DOUBLE HEX"]),
    ("BOLT, HEX",        ["NUT, PLAIN, HEX", "NUT, SELF-LOCKING, HEX"]),
    # T-HEAD 볼트의 짝은 확인되지 않았다. 추정하지 않는다(HANDOFF '모르면 모른다고').
    ("BOLT, T-HEAD",     []),
    ("STUD, RING LOCKED, SERRATED", ["NUT, PLAIN, HEX", "NUT, SELF-LOCKING, HEX"]),
    ("STUD, KEY LOCKED",            ["NUT, PLAIN, HEX", "NUT, SELF-LOCKING, HEX"]),
])

# 되물을 때 사람에게 보여줄 구분 설명
NUT_CHOICE_HINT = {
    "NUT, PLAIN, HEX": "일반(락킹 없음) — 별도 풀림방지 수단이 있을 때",
    "NUT, SELF-LOCKING, HEX": "셀프락킹 — 진동부 기본",
    "NUT, SELF-LOCKING, DOUBLE HEX": "셀프락킹 더블헥스",
}


def nut_categories_for(bolt_mid):
    """볼트 중분류 -> (너트 중분류 후보, 안내문).

    후보가 0개면 '짝 없음'을 뜻한다 — 억지로 추천하지 않는다.
    후보가 2개 이상이면 **사용자에게 되물어야 한다**. 자동으로 하나를 고르지 않는다.
    """
    key = norm(bolt_mid)
    for k, v in NUT_FOR_BOLT.items():
        if norm(k) == key:
            if not v:
                return [], f"'{bolt_mid}' 의 짝이 되는 너트 중분류가 확인되지 않았습니다."
            if len(v) == 1:
                return list(v), ""
            opts = "; ".join(f"{n}({NUT_CHOICE_HINT.get(n, '')})" for n in v)
            return list(v), f"너트 종류를 정해야 합니다 -> {opts}"
    return [], f"'{bolt_mid}' 는 매핑표에 없는 중분류입니다. NUT_FOR_BOLT 에 추가하세요."


# ---------------------------------------------------------------- 직경 파싱


def parse_diameter(pumyeong):
    """품명에서 직경을 뽑는다. 'BOLT, DOUBLE HEX, 0.164' -> 0.164. 없으면 None.

    품명의 마지막 조각이 숫자면 그것이 직경이다. 정수만 있는 조각(예: 사이즈 코드)은
    직경으로 보지 않는다 — 표준품 직경은 항상 소수 인치 표기다.
    """
    if not pumyeong:
        return None
    tail = str(pumyeong).split(",")[-1].strip()
    m = re.fullmatch(r"(\d*\.\d+)", tail)
    return float(m.group(1)) if m else None


def diameter_matches(a, b, tol=0.0005):
    """직경 비교. .164 와 0.1640 을 같게 본다."""
    if a is None or b is None:
        return False
    return abs(a - b) <= tol


# ---------------------------------------------------------------- 엑셀/CSV 로딩

# 엑셀 헤더 이름 -> 내부 키. 실제 파일 헤더가 다르면 load_catalog(columns=...) 로 넘긴다.
DEFAULT_COLUMNS = {
    "대분류": "major", "중분류": "mid", "도면번호": "drawing",
    "부품번호": "part", "품명": "name", "소재": "material", "비고": "note",
}


class Item(dict):
    """카탈로그 한 행. dict 그대로 쓰되 자주 쓰는 필드는 속성으로도 꺼낸다."""

    @property
    def major(self):
        return self.get("major", "")

    @property
    def mid(self):
        return self.get("mid", "")

    @property
    def drawing(self):
        return self.get("drawing", "")

    @property
    def part(self):
        return self.get("part", "")

    @property
    def name(self):
        return self.get("name", "")

    @property
    def diameter(self):
        return self.get("diameter")


def _rows_from_xlsx(path, sheet=None):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError(
            "xlsx 를 읽으려면 openpyxl 이 필요합니다. 사내 미러에 없으면 엑셀에서 "
            "CSV 로 저장해 그 파일을 쓰세요(load_catalog 가 .csv 도 읽습니다)."
        )
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def _rows_from_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return [r for r in csv.reader(f)]


def load_catalog(path, sheet=None, columns=None):
    """엑셀(.xlsx) 또는 CSV 표준품 목록을 읽어 Item 리스트로.

    첫 행을 헤더로 본다. 알아본 컬럼만 담고 나머지는 무시한다.
    품명에서 직경을 파싱해 'diameter' 로 넣어 둔다(엑셀에 직경 컬럼이 없다).
    """
    if not os.path.exists(path):
        raise RuntimeError(f"표준품 목록 파일이 없습니다: {path}")
    raw = _rows_from_csv(path) if path.lower().endswith(".csv") else _rows_from_xlsx(path, sheet)
    if not raw:
        raise RuntimeError(f"빈 파일입니다: {path}")

    colmap = dict(DEFAULT_COLUMNS)
    if columns:
        colmap.update(columns)
    header = [norm(c) for c in raw[0]]
    idx = {}
    for label, key in colmap.items():
        n = norm(label)
        if n in header:
            idx[key] = header.index(n)
    if "mid" not in idx or "drawing" not in idx:
        raise RuntimeError(
            f"필수 컬럼을 찾지 못했습니다. 읽은 헤더: {raw[0]}\n"
            f"기대: {list(colmap)} — 다르면 load_catalog(columns={{...}}) 로 알려주세요."
        )

    items = []
    for r in raw[1:]:
        if not any(c not in (None, "") for c in r):
            continue  # 빈 행
        item = Item()
        for key, i in idx.items():
            item[key] = str(r[i]).strip() if i < len(r) and r[i] is not None else ""
        item["diameter"] = parse_diameter(item.get("name"))
        items.append(item)
    return items


def validate_catalog(items):
    """엑셀 값이 아는 중분류인지 점검한다. 반환: 사람이 볼 알림 리스트.

    모르는 중분류를 조용히 넘기면 매핑에서 빠져 '짝 없음'으로 **잘못** 결론난다.
    그래서 판단이 아니라 보고를 한다.
    """
    notes = []
    unknown = OrderedDict()
    for it in items:
        if it.mid and norm(it.mid) not in _KNOWN:
            unknown.setdefault(it.mid, 0)
            unknown[it.mid] += 1
    for mid, n in unknown.items():
        notes.append(f"모르는 중분류 '{mid}' ({n}행) — KNOWN_MIDS/NUT_FOR_BOLT 확인 필요")

    no_dia = sum(1 for it in items if it.major and it.diameter is None)
    if no_dia:
        notes.append(f"품명에서 직경을 못 읽은 행 {no_dia}개 (직경 조회는 그 행을 못 찾는다)")
    return notes


# ---------------------------------------------------------------- 조회


def find_series(items, major=None, mid=None, diameter=None):
    """조건에 맞는 계열(도면번호) 목록. 반환: [(도면번호, 품명, 직경, 행수), ...]

    한 계열 = 도면번호 하나. dash 는 여기서 정하지 않는다(스펙 판독의 몫).
    """
    hit = OrderedDict()
    for it in items:
        if major and norm(it.major) != norm(major):
            continue
        if mid and norm(it.mid) != norm(mid):
            continue
        if diameter is not None and not diameter_matches(it.diameter, diameter):
            continue
        key = it.drawing or it.name
        if key not in hit:
            hit[key] = {"drawing": it.drawing, "name": it.name,
                        "diameter": it.diameter, "n": 0}
        hit[key]["n"] += 1
    return [(v["drawing"], v["name"], v["diameter"], v["n"]) for v in hit.values()]


def get_spec_path(items_or_name, spec_dir, ext=".pdf"):
    """스펙 PDF 경로를 찾는다. 없으면 None.

    ⚠ 문서 두 곳이 서로 다르게 적혀 있다 — HANDOFF.md 는 '품명 = 파일명',
    CLAUDE.md 는 '도면번호 = 파일명'. README 의 실행 예시는 품명 쪽이다
    ("BOLT, DOUBLE HEX, 0.164.pdf"). 어느 쪽인지 확정될 때까지 **둘 다 시도**하고,
    무엇으로 찾았는지는 호출부가 알 수 있게 경로를 그대로 돌려준다.
    """
    name = items_or_name.name if isinstance(items_or_name, Item) else str(items_or_name)
    drawing = items_or_name.drawing if isinstance(items_or_name, Item) else ""
    for cand in (name, drawing):
        if not cand:
            continue
        p = os.path.join(spec_dir, cand + ext)
        if os.path.exists(p):
            return p
    return None


def summarize(items):
    """목록 개요 한 줄 (사내망에서 눈으로 확인용)."""
    majors = OrderedDict()
    for it in items:
        if it.major:
            majors[it.major] = majors.get(it.major, 0) + 1
    top = ", ".join(f"{k} {v}" for k, v in list(majors.items())[:6])
    return f"부품번호 {len(items)}행 / 대분류 {len(majors)}종 ({top}{' ...' if len(majors) > 6 else ''})"
